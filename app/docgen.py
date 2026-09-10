"""Geracao dos documentos (declaracoes) em texto renderizado e arquivo .docx.

O layout reproduz o padrao usado na planilha original:
    cabecalho com logo (esquerda) + endereco da empresa (direita)
    Data: dd/mm/aaaa (direita)
    destinatario / assunto / modelos
    corpo da declaracao
    dados do responsavel + assinatura
"""
from datetime import date
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from flask import current_app
from jinja2 import Environment, BaseLoader

from .models import Empresa
from .utils import caminho_absoluto, data_br, data_extenso, pasta_processo, slugify

CINZA = RGBColor(0x47, 0x4B, 0x4F)
TEXTO = RGBColor(0x55, 0x55, 0x55)
FONTE = "Calibri"

SEPARADOR_EN = "---EN---"          # divide o corpo em duas colunas (portugues | ingles)
MARCA_REQUISITOS = "---REQUISITOS---"   # onde entra a tabela C / NA da seguranca cibernetica

_env = Environment(loader=BaseLoader(), autoescape=False, trim_blocks=True, lstrip_blocks=True)
_env.filters["data_extenso"] = data_extenso
_env.filters["data_br"] = data_br


def marcar(valor):
    """Caixa de selecao usada nas declaracoes de rastreabilidade / manutencao."""
    return "☒" if valor else "☐"


def contexto(processo, signatario=None, cidade=None, data_doc=None, extras=None):
    """Variaveis disponiveis nos templates de documento."""
    empresa = db_empresa()
    produto = processo.produto
    sig = signatario or processo.signatario
    ctx = {
        "empresa": empresa,
        "processo": processo,
        "produto": produto,
        "fabricante": produto.fabricante if produto else None,
        "ocd": processo.ocd,
        "laboratorio": processo.laboratorio,
        "signatario": sig,
        "similares": list(processo.similares),
        "amostras": list(processo.amostras),
        "hoje": data_doc or date.today(),
        "data_extenso": data_extenso(data_doc or date.today()),
        "data": data_br(data_doc or date.today()),
        "cidade": cidade or (empresa.cidade_assinatura if empresa else "")
                  or (empresa.cidade if empresa else ""),
        "modelo": produto.modelo if produto else "",
        "modelos": ", ".join(processo.modelos_declarados),
        "modelos_similares": ", ".join(s.modelo for s in processo.similares),
        "marcar": marcar,
        "rastreabilidade": processo.rastreabilidade_itens,
    }
    if extras:
        ctx.update(extras)
    return ctx


def db_empresa():
    return Empresa.query.get(1)


def renderizar(texto, ctx):
    try:
        return _env.from_string(texto or "").render(**ctx).strip()
    except Exception as exc:  # um template com erro nao pode derrubar a tela
        return f"[ERRO NO TEMPLATE: {exc}]\n\n{texto or ''}"


def dividir_bilingue(corpo):
    """Devolve (portugues, ingles) quando o corpo tem o separador ---EN---."""
    if SEPARADOR_EN in (corpo or ""):
        pt, en = corpo.split(SEPARADOR_EN, 1)
        return pt.strip(), en.strip()
    return (corpo or "").strip(), None


def blocos(texto):
    return [b.rstrip() for b in (texto or "").split("\n\n") if b.strip()]


def partir_requisitos(corpo):
    """Devolve (antes, depois) do marcador ---REQUISITOS---, ou (corpo, None)."""
    if MARCA_REQUISITOS in (corpo or ""):
        antes, depois = corpo.split(MARCA_REQUISITOS, 1)
        return antes.strip(), depois.strip()
    return (corpo or "").strip(), None


# ----------------------------------------------------------------- docx

def _sem_bordas(tabela):
    tbl = tabela._tbl
    props = tbl.tblPr
    for filho in props.findall(qn("w:tblBorders")):
        props.remove(filho)
    return tabela


def _remover_paragrafo_vazio(celula):
    """Tira o paragrafo vazio que o Word cria junto com a celula."""
    if len(celula.paragraphs) > 1 and not celula.paragraphs[0].text.strip():
        elemento = celula.paragraphs[0]._element
        elemento.getparent().remove(elemento)


def _run(p, texto, tam=11, negrito=False, italico=False, cor=TEXTO):
    r = p.add_run(texto)
    r.bold = negrito
    r.italic = italico
    r.font.size = Pt(tam)
    r.font.name = FONTE
    r.font.color.rgb = cor
    return r


def _par(container, texto="", tam=11, negrito=False, alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
         depois=6, cor=TEXTO, italico=False, entrelinha=1.15):
    p = container.add_paragraph()
    p.alignment = alinhamento
    p.paragraph_format.space_after = Pt(depois)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = entrelinha
    linhas = str(texto).split("\n")
    for i, linha in enumerate(linhas):
        if i:
            p.add_run().add_break()
        _run(p, linha, tam=tam, negrito=negrito, italico=italico, cor=cor)
    return p


def _bloco_corpo(container, texto, tam=11):
    """Um bloco de texto: varias linhas viram quebras; prosa longa fica justificada."""
    linhas = [l for l in texto.split("\n")]
    uma_linha = len([l for l in linhas if l.strip()]) == 1
    marcador = texto.lstrip().startswith(("•", "-", "*", "☐", "☒", "1)", "2)", "3)"))
    if uma_linha and not marcador and len(texto) > 110:
        alinhamento = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        alinhamento = WD_ALIGN_PARAGRAPH.LEFT
    return _par(container, texto, tam=tam, alinhamento=alinhamento, depois=8)


def _cabecalho(sec, empresa):
    """Logo a esquerda, endereco a direita."""
    p0 = sec.header.paragraphs[0]
    p0.text = ""
    tabela = sec.header.add_table(rows=1, cols=2, width=Cm(16.5))
    _sem_bordas(tabela)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    esq, dir_ = tabela.rows[0].cells
    esq.width = Cm(6.0)
    dir_.width = Cm(10.5)

    logo = None
    if empresa and empresa.logo_arquivo:
        caminho = caminho_absoluto(empresa.logo_arquivo)
        if caminho.is_file():
            logo = caminho
    pl = esq.paragraphs[0]
    pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if logo:
        try:
            pl.add_run().add_picture(str(logo), width=Cm(4.0))
        except Exception:
            _run(pl, empresa.nome_fantasia or empresa.razao_social, tam=12, negrito=True, cor=CINZA)
    elif empresa:
        _run(pl, empresa.nome_fantasia or empresa.razao_social, tam=12, negrito=True, cor=CINZA)

    pd = dir_.paragraphs[0]
    pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pd.paragraph_format.space_after = Pt(0)
    if empresa:
        for i, linha in enumerate(empresa.cabecalho):
            if i:
                pd.add_run().add_break()
            _run(pd, linha, tam=8.5, cor=TEXTO)


def _tabela_requisitos(doc, processo):
    """Tabela C / NA dos requisitos de seguranca cibernetica."""
    from . import requisitos_ciber
    if not processo.requisitos_ciber:
        _par(doc, "[o formulário de requisitos ainda não foi preenchido]", italico=True,
             cor=TEXTO)
        return

    for titulo, secoes in requisitos_ciber.agrupado(processo.requisitos_ciber):
        if titulo:
            _par(doc, titulo, negrito=True, tam=10, depois=6, cor=CINZA)
        for secao, itens in secoes:
            if secao:
                _par(doc, secao, negrito=True, tam=10, depois=4, cor=CINZA)
            tabela = doc.add_table(rows=0, cols=2)
            tabela.style = "Table Grid"
            for r in itens:
                celulas = tabela.add_row().cells
                celulas[0].width = Cm(1.4)
                celulas[1].width = Cm(14.6)
                p_situacao = celulas[0].paragraphs[0]
                p_situacao.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _run(p_situacao, r.situacao, tam=10, negrito=True, cor=CINZA)

                celulas[1].paragraphs[0].text = ""
                p_texto = celulas[1].paragraphs[0]
                _run(p_texto, f"{r.codigo} ", tam=10, negrito=True, cor=CINZA)
                _run(p_texto, r.texto, tam=10)
                if r.situacao == "NA":
                    p_just = celulas[1].add_paragraph()
                    p_just.paragraph_format.space_before = Pt(2)
                    p_just.paragraph_format.space_after = Pt(0)
                    _run(p_just, "Justificativa: ", tam=9, negrito=True, cor=CINZA)
                    _run(p_just, r.justificativa or "—", tam=9, italico=True)
            _par(doc, "", depois=6)


def gerar_docx(documento):
    """Monta o .docx e devolve o caminho relativo salvo em instance/arquivos."""
    processo = documento.processo
    empresa = db_empresa()
    sig = documento.signatario or processo.signatario

    doc = DocxDocument()
    estilo = doc.styles["Normal"]
    estilo.font.name = FONTE
    estilo.font.size = Pt(11)

    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    _cabecalho(sec, empresa)

    # ---- data
    _par(doc, f"Data: {data_br(documento.data_documento)}",
         alinhamento=WD_ALIGN_PARAGRAPH.RIGHT, depois=14, tam=10.5)

    pt, en = dividir_bilingue(documento.corpo)

    if en:
        # documento bilingue: portugues | ingles, lado a lado
        tabela = doc.add_table(rows=1, cols=2)
        _sem_bordas(tabela)
        col_pt, col_en = tabela.rows[0].cells
        col_pt.width = Cm(8.0)
        col_en.width = Cm(8.0)
        for bloco in blocos(pt):
            _bloco_corpo(col_pt, bloco, tam=10)
        for bloco in blocos(en):
            _bloco_corpo(col_en, bloco, tam=10)
        for cel in (col_pt, col_en):
            _remover_paragrafo_vazio(cel)
        _par(doc, "", depois=4)
    else:
        antes, depois = partir_requisitos(pt)
        for bloco in blocos(antes):
            _bloco_corpo(doc, bloco)
        if depois is not None:
            _tabela_requisitos(doc, processo)
            for bloco in blocos(depois):
                _bloco_corpo(doc, bloco)

    # ---- tabela de modelos similares
    if documento.tipo == "similaridade" and processo.similares:
        tabela = doc.add_table(rows=1, cols=3)
        tabela.style = "Table Grid"
        for i, texto in enumerate(("Modelo", "Nome comercial",
                                   "Diferenças em relação ao modelo base")):
            cel = tabela.rows[0].cells[i]
            cel.paragraphs[0].text = ""
            _run(cel.paragraphs[0], texto, tam=9.5, negrito=True, cor=CINZA)
        for s in processo.similares:
            celulas = tabela.add_row().cells
            for i, valor in enumerate((s.modelo, s.nome_comercial or "-", s.diferencas or "-")):
                celulas[i].paragraphs[0].text = ""
                _run(celulas[i].paragraphs[0], valor, tam=9.5)
        _par(doc, "", depois=8)

    # ---- dados do responsavel (alguns formularios ja os trazem no proprio texto)
    _par(doc, "", depois=12)
    mostrar_responsavel = not (documento.template and not documento.template.bloco_responsavel)
    if sig and mostrar_responsavel:
        _par(doc, f"Nome do responsável: {sig.nome}", depois=0, tam=10.5)
        if sig.cargo:
            _par(doc, f"Cargo: {sig.cargo}", depois=0, tam=10.5)
        if sig.email:
            _par(doc, f"E-mail: {sig.email}", depois=0, tam=10.5)
        if sig.email_alternativo:
            _par(doc, f"E-mail: {sig.email_alternativo}", depois=0, tam=10.5)
        if sig.telefone:
            _par(doc, f"Telefone: {sig.telefone}", depois=0, tam=10.5)

    # ---- local, data e assinatura
    _par(doc, "", depois=10)
    local_data = f"{documento.cidade or ''}, {data_extenso(documento.data_documento)}."
    _par(doc, local_data.lstrip(", "), alinhamento=WD_ALIGN_PARAGRAPH.LEFT, depois=18, tam=10.5)

    _par(doc, "Assinatura:", depois=6, tam=10.5)
    if sig and sig.assinatura_arquivo:
        caminho = caminho_absoluto(sig.assinatura_arquivo)
        if caminho.is_file():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            try:
                p.add_run().add_picture(str(caminho), width=Cm(4.5))
            except Exception:
                pass
    _par(doc, "_" * 41, depois=2, cor=CINZA)
    if sig:
        _par(doc, sig.nome, negrito=True, depois=0, cor=CINZA, tam=10.5)

    # ---- rodape com rastreabilidade interna do sistema
    rod = sec.footer.paragraphs[0]
    rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emitido = documento.criado_em.strftime("%d/%m/%Y %H:%M") if documento.criado_em else ""
    r = rod.add_run(f"{processo.numero}  |  {documento.titulo}  |  v{documento.versao}"
                    f"  |  emitido em {emitido}")
    r.font.size = Pt(7)
    r.font.name = FONTE
    r.font.color.rgb = TEXTO

    destino = pasta_processo(processo, "documentos")
    nome = f"{slugify(processo.numero)}-{slugify(documento.tipo)}-v{documento.versao}.docx"
    caminho = destino / nome
    doc.save(caminho)
    raiz = Path(current_app.config["UPLOAD_DIR"])
    return caminho.relative_to(raiz).as_posix()
